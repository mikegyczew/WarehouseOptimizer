try:
    from app.optimizer.packing import Box, optimize_packing
except ModuleNotFoundError:  # pragma: no cover
    from optimizer.packing import Box, optimize_packing

from io import BytesIO
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel


class Package(BaseModel):
    name: str
    length: float
    width: float
    height: float
    quantity: int


class OptimizationRequest(BaseModel):
    room_length: float
    room_width: float
    room_height: float
    packages: list[Package]


class ExcelImportService:
    @staticmethod
    def _to_float(value):
        if value is None:
            return 0.0
        if isinstance(value, str):
            value = value.strip().replace(",", ".")
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _normalize(value):
        if value is None:
            return ""
        return str(value).strip().lower().replace(" ", "_")

    @staticmethod
    def build_template_workbook() -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Warehouse"
        sheet.append(["room_length", "room_width", "room_height"])
        sheet.append([1000, 800, 300])
        sheet.append([])
        sheet.append(["name", "length", "width", "height", "quantity"])
        sheet.append(["Karton A", 120, 60, 50, 12])
        sheet.append(["Karton B", 80, 50, 40, 20])
        sheet.append(["Paczka C", 60, 40, 30, 35])
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = max_length + 2
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer

    @staticmethod
    def parse_excel(file_bytes: bytes) -> dict:
        workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            raise ValueError("Excel jest pusty.")

        room = {}
        for index, row in enumerate(rows):
            normalized = [ExcelImportService._normalize(value) for value in row]
            label_positions = {
                key: position
                for position, key in enumerate(normalized)
                if key in {"room_length", "room_width", "room_height"}
            }
            if not label_positions:
                continue
            next_row = rows[index + 1] if index + 1 < len(rows) else []
            for key, position in label_positions.items():
                if position < len(next_row) and next_row[position] not in (None, ""):
                    room[key] = ExcelImportService._to_float(next_row[position])
            if room:
                break

        if not room:
            room_values = []
            for row in rows[:3]:
                room_values.extend(
                    ExcelImportService._to_float(value)
                    for value in row[:3]
                    if value not in (None, "")
                )
            if len(room_values) >= 3:
                room = {
                    "room_length": room_values[0],
                    "room_width": room_values[1],
                    "room_height": room_values[2],
                }

        if not room:
            raise ValueError("Nie znaleziono wymiarów magazynu w pliku Excel.")

        package_rows = []
        header_positions = {}
        for index, row in enumerate(rows):
            normalized = [ExcelImportService._normalize(value) for value in row]
            if any(value in {"name", "length", "width", "height", "quantity"} for value in normalized):
                header_positions = {
                    key: position
                    for position, key in enumerate(normalized)
                    if key in {"name", "length", "width", "height", "quantity"}
                }
                for package_index in range(index + 1, len(rows)):
                    package_row = rows[package_index]
                    if not package_row or all(value in (None, "") for value in package_row):
                        continue
                    name_value = package_row[header_positions.get("name", 0)] if "name" in header_positions and header_positions["name"] < len(package_row) else ""
                    length_value = package_row[header_positions.get("length", 1)] if "length" in header_positions and header_positions["length"] < len(package_row) else 0
                    width_value = package_row[header_positions.get("width", 2)] if "width" in header_positions and header_positions["width"] < len(package_row) else 0
                    height_value = package_row[header_positions.get("height", 3)] if "height" in header_positions and header_positions["height"] < len(package_row) else 0
                    quantity_value = package_row[header_positions.get("quantity", 4)] if "quantity" in header_positions and header_positions["quantity"] < len(package_row) else 1
                    if name_value in (None, ""):
                        continue
                    package_rows.append(
                        {
                            "name": str(name_value).strip(),
                            "length": ExcelImportService._to_float(length_value),
                            "width": ExcelImportService._to_float(width_value),
                            "height": ExcelImportService._to_float(height_value),
                            "quantity": int(ExcelImportService._to_float(quantity_value) or 1),
                        }
                    )
                break

        if not package_rows:
            package_rows = []

        return {
            "room_length": room.get("room_length") or 0,
            "room_width": room.get("room_width") or 0,
            "room_height": room.get("room_height") or 0,
            "packages": package_rows,
        }


class OptimizationService:
    @staticmethod
    def build_boxes(packages: list[Package]) -> list[Box]:
        return [
            Box(
                name=package.name,
                length=package.length,
                width=package.width,
                height=package.height,
                quantity=package.quantity,
            )
            for package in packages
        ]

    @staticmethod
    def calculate_result(data: OptimizationRequest) -> dict:
        boxes = OptimizationService.build_boxes(data.packages)
        placed, not_placed = optimize_packing(
            room_length=data.room_length,
            room_width=data.room_width,
            room_height=data.room_height,
            boxes=boxes,
        )

        room_volume = (
            data.room_length * data.room_width * data.room_height / 1_000_000
        )
        placed_volume = sum(
            box.length * box.width * box.height / 1_000_000 for box in placed
        )
        utilization = (
            (placed_volume / room_volume) * 100 if room_volume > 0 else 0
        )

        return {
            "success": True,
            "room": {
                "length": data.room_length,
                "width": data.room_width,
                "height": data.room_height,
                "volume": round(room_volume, 3),
            },
            "optimization": {
                "placed": len(placed),
                "not_placed": len(not_placed),
                "utilization": round(utilization, 2),
            },
            "placed_boxes": [
                {
                    "name": box.name,
                    "length": box.length,
                    "width": box.width,
                    "height": box.height,
                    "x": box.x,
                    "y": box.y,
                    "z": box.z,
                }
                for box in placed
            ],
            "not_placed_boxes": not_placed,
        }


class WarehouseAppFactory:
    @staticmethod
    def create_app() -> FastAPI:
        app = FastAPI(title="WarehouseOptimizer")
        base_dir = Path(__file__).resolve().parent

        app.mount(
            "/static",
            StaticFiles(directory=str(base_dir / "static")),
            name="static",
        )
        templates = Jinja2Templates(directory=str(base_dir / "templates"))

        @app.get("/", response_class=HTMLResponse)
        async def home(request: Request):
            return templates.TemplateResponse(
                request=request,
                name="index.html",
                context={},
            )

        @app.get("/api/excel-template")
        async def excel_template():
            template = ExcelImportService.build_template_workbook()
            return StreamingResponse(
                template,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": 'attachment; filename="warehouse_template.xlsx"'},
            )

        @app.post("/api/excel-import")
        async def excel_import(file: UploadFile = File(...)):
            if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm", ".xls")):
                raise HTTPException(status_code=400, detail="Dozwolone są tylko pliki Excel (.xlsx, .xlsm, .xls).")
            try:
                data = ExcelImportService.parse_excel(await file.read())
            except Exception as exc:  # pragma: no cover - runtime validation path
                raise HTTPException(status_code=400, detail=str(exc)) from exc
            return {"success": True, "data": data}

        @app.post("/api/optimize")
        async def optimize(data: OptimizationRequest):
            return OptimizationService.calculate_result(data)

        return app


app = WarehouseAppFactory.create_app()